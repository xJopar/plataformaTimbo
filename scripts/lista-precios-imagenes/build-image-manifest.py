"""
Fase D del pipeline de imagenes de Lista de Precios: genera el manifiesto Stock -> prefijo de
carpeta en el bucket, a partir de:
  1. El arbol de carpetas ya organizado/renombrado a mano (imagenes_lista_precios/), que es un
     espejo exacto de lo subido al bucket de Railway.
  2. V_ListaPreciosTableau2.xlsx (Stock, Marca, Modelo, Config., Susp.).

El arbol de carpetas ya no es 100% derivable por formula (el usuario reorganizo a mano varias
carpetas), asi que esto NO vuelve a bajar ni resolver permalinks: solo empareja carpetas ya
existentes con filas de la planilla, usando esta precedencia (mas especifico gana):

  1. Una subcarpeta en cualquier nivel bajo MARCA/MODELO cuyo nombre sea exactamente el Stock
     de la fila (override por unidad).
  2. Una subcarpeta bajo MARCA/MODELO cuyo nombre (saneado) matchea el valor de Config. y/o
     Susp. de la fila (override por variante/config).
  3. Los archivos que esten directamente en MARCA/MODELO/ (default de todo el modelo).
  4. Sin match -> la fila queda sin imagenes.

Salida:
  - vehicle-image-folders.json  { "<STOCK>": "MARCA/MODELO/.../" }
  - _manifest_reporte.csv       una fila por Stock con el prefijo resuelto y el metodo usado

Uso:
  python build-image-manifest.py <V_ListaPreciosTableau2.xlsx> <carpetaImagenes> <salida.json> <salida.csv>
"""

import csv
import json
import os
import re
import sys
import unicodedata

import openpyxl


def strip_accents(value):
    normalized = unicodedata.normalize("NFKD", value)
    return "".join(c for c in normalized if not unicodedata.combining(c))


def sanitize(value):
    cleaned = strip_accents((value or "").strip().upper())
    cleaned = re.sub(r"[^A-Z0-9]+", "_", cleaned)
    cleaned = re.sub(r"_+", "_", cleaned).strip("_")
    return cleaned


def digits_only(value):
    return re.sub(r"[^0-9]", "", value or "")


def fields_match(folder_key, field_value):
    """True si el nombre saneado de la carpeta matchea el valor de campo, ya sea por
    contención de texto o, cuando ambos son mayormente numéricos (ej. "14-50mts" vs
    "14,5 MTS."), por contención de sus dígitos."""
    if not folder_key or not field_value:
        return False
    if folder_key == field_value or folder_key in field_value or field_value in folder_key:
        return True
    folder_digits = digits_only(folder_key)
    field_digits = digits_only(field_value)
    if folder_digits and field_digits:
        return folder_digits in field_digits or field_digits in folder_digits
    return False


STOCK_PATTERN = re.compile(r"^C\d+$", re.IGNORECASE)


class Node:
    def __init__(self, name, path_segments):
        self.name = name
        self.path_segments = path_segments
        self.children = {}
        self.has_own_files = False

    def prefix(self):
        return "/".join(self.path_segments)


def build_tree(root_dir):
    root = Node("", [])
    for current_dir, dirnames, filenames in os.walk(root_dir):
        rel = os.path.relpath(current_dir, root_dir)
        segments = [] if rel == "." else rel.split(os.sep)
        node = root
        for segment in segments:
            node = node.children.setdefault(segment, Node(segment, node.path_segments + [segment]))
        real_files = [f for f in filenames if not f.startswith("_")]
        if real_files:
            node.has_own_files = True
    return root


def find_stock_node(node, stock_upper):
    if node.name.upper() == stock_upper:
        return node
    for child in node.children.values():
        found = find_stock_node(child, stock_upper)
        if found is not None:
            return found
    return None


def resolve_prefix(modelo_node, row):
    """Devuelve (prefix_or_None, metodo) para una fila, buscando dentro de modelo_node."""
    stock_upper = (row["stock"] or "").upper()

    # 1) Override por Stock, en cualquier profundidad.
    if stock_upper:
        stock_node = find_stock_node(modelo_node, stock_upper)
        if stock_node is not None and stock_node.has_own_files:
            return stock_node.prefix(), "stock"

    # 2) Override por subcarpeta de config/susp, recursivo, el mas profundo que matchea gana.
    # Config. se prueba antes que Susp. en cada nivel: "TUMBA 18M3" con Susp="6X4" debe caer en
    # la carpeta "Tumba" (matchea Config), no en "chasis_6x4" (matchea Susp), porque el tipo de
    # carrocería (Config) es una distinción mas fuerte que la config de ejes (Susp).
    config_field = sanitize(row.get("config", ""))
    susp_field = sanitize(row.get("susp", ""))

    best = None  # (depth, prefix)

    def matching_children(node, field_value):
        for child_name, child in node.children.items():
            if STOCK_PATTERN.match(child_name):
                continue  # las carpetas de stock ya se resolvieron arriba
            if fields_match(sanitize(child_name), field_value):
                yield child

    def walk(node, remaining_config, remaining_susp):
        nonlocal best
        matched_any = False
        if remaining_config:
            for child in matching_children(node, remaining_config):
                matched_any = True
                register(child)
                walk(child, "", remaining_susp)
        if not matched_any and remaining_susp:
            for child in matching_children(node, remaining_susp):
                register(child)
                walk(child, remaining_config, "")

    def register(child):
        nonlocal best
        if child.has_own_files:
            depth = len(child.path_segments)
            if best is None or depth > best[0]:
                best = (depth, child.prefix())

    walk(modelo_node, config_field, susp_field)
    if best is not None:
        return best[1], "config_susp"

    # 3) Default de modelo completo.
    if modelo_node.has_own_files:
        return modelo_node.prefix(), "modelo_default"

    return None, "sin_match"


def main():
    if len(sys.argv) != 5:
        print(
            "Uso: python build-image-manifest.py <V_ListaPreciosTableau2.xlsx> "
            "<carpetaImagenes> <salida.json> <salida.csv>"
        )
        sys.exit(1)

    excel_path, images_dir, out_json, out_csv = sys.argv[1:5]

    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_index = {h: i for i, h in enumerate(headers)}

    def get(row_cells, col_name):
        idx = col_index.get(col_name)
        if idx is None or idx >= len(row_cells):
            return ""
        value = row_cells[idx]
        return str(value).strip() if value is not None else ""

    rows = []
    for row_cells in ws.iter_rows(min_row=2, values_only=True):
        stock = get(row_cells, "Stock")
        if not stock:
            continue
        rows.append(
            {
                "stock": stock,
                "marca": get(row_cells, "Marca"),
                "modelo": get(row_cells, "Modelo"),
                "config": get(row_cells, "Config."),
                "susp": get(row_cells, "Susp."),
            }
        )

    tree = build_tree(images_dir)

    manifest = {}
    report_rows = []
    counts = {"stock": 0, "config_susp": 0, "modelo_default": 0, "sin_match": 0}

    for row in rows:
        marca_node = tree.children.get(sanitize_lookup(tree, row["marca"]))
        prefix = None
        metodo = "sin_match"
        if marca_node is not None:
            modelo_node = marca_node.children.get(sanitize_lookup(marca_node, row["modelo"]))
            if modelo_node is not None:
                prefix, metodo = resolve_prefix(modelo_node, row)

        counts[metodo] += 1
        if prefix is not None:
            manifest[row["stock"]] = prefix

        report_rows.append(
            {
                "stock": row["stock"],
                "marca": row["marca"],
                "modelo": row["modelo"],
                "config": row["config"],
                "susp": row["susp"],
                "metodo": metodo,
                "prefijoResuelto": prefix or "",
            }
        )

    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2, sort_keys=True)

    with open(out_csv, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f, fieldnames=["stock", "marca", "modelo", "config", "susp", "metodo", "prefijoResuelto"]
        )
        writer.writeheader()
        writer.writerows(report_rows)

    print(f"Filas de entrada: {len(rows)}")
    print(f"Resueltas por Stock (override unidad): {counts['stock']}")
    print(f"Resueltas por Config/Susp (override variante): {counts['config_susp']}")
    print(f"Resueltas por default de modelo: {counts['modelo_default']}")
    print(f"Sin match: {counts['sin_match']}")
    print(f"Manifiesto: {out_json}")
    print(f"Reporte: {out_csv}")


def sanitize_lookup(node, raw_name):
    """Busca, entre los hijos directos de `node`, cual nombre real sanea igual a raw_name."""
    target = sanitize(raw_name)
    for child_name in node.children:
        if sanitize(child_name) == target:
            return child_name
    return None


if __name__ == "__main__":
    main()
