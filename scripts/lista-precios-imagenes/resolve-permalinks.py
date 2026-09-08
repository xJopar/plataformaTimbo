"""
Fase A del plan de imagenes de Lista de Precios.

Reproduce exactamente esta logica de negocio (provista por el usuario) para resolver, por cada
Stock de V_ListaPreciosTableau2, el permalink de timbo.com.py que tiene la foto a usar:

    SELECT V.*, COALESCE(E1.permalink, E2.permalink, E3.permalink) AS permalink
    FROM V_ListaPreciosTableau2 V
    LEFT JOIN ELICE_URLS E1 ON V.Stock = E1.sku
    LEFT JOIN ELICE_URLS E2 ON V.CodGrupoUnidad = E2.sku AND E1.permalink IS NULL
    LEFT JOIN ELICE_URLS E3 ON <CASE de normalizacion de CodGrupoUnidad> = E3.sku
                            AND E1.permalink IS NULL AND E2.permalink IS NULL

E1 = foto propia de la unidad (match exacto por Stock).
E2 = foto default curada para el CodGrupoUnidad exacto de la unidad.
E3 = foto default curada para el "grupo padre" de ese CodGrupoUnidad (via el CASE).

Las unidades que no matchean en ninguna capa (CodGrupoUnidad sin categorizar en SAP, o marca sin
grupo curado en ELICE) quedan con permalink vacio a proposito: no se les asigna una imagen
"adivinada" por busqueda semantica.

Uso:
    python resolve-permalinks.py <V_ListaPreciosTableau2.xlsx> <ELICE_URLS.xlsx> <salida.json>
"""

import json
import sys
from collections import Counter

import openpyxl

# CASE de normalizacion, tal cual la definicion de negocio.
GROUP_NORMALIZATION = {}
for members, target in [
    (["L2", "L3", "L4"], "L1"),
    (["L5", "L6", "L7", "L8"], "L5"),
    (["L10", "L11"], "L10"),
    (["L20", "L21", "L22", "L23", "L24"], "L23"),
    (["F5", "F14"], "F5"),
    (["F9", "F10"], "F10"),
    (["F11", "F12"], "F11"),
    (["ST11", "ST12", "ST8"], "ST11"),
    (["ST10"], "ST9"),
    (["ST7"], "ST6"),
    (["ST3"], "ST2"),
    (["ST15"], "ST14"),
    (["SN2"], "SN1"),
    (["SN4"], "SN3"),
    (["SN16"], "SN15"),
]:
    for member in members:
        GROUP_NORMALIZATION.setdefault(member, target)


def normalize_group(cod_grupo_unidad):
    return GROUP_NORMALIZATION.get(cod_grupo_unidad, cod_grupo_unidad)


def load_elice_sku_index(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    sku_to_permalink = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        sku, _name, permalink = row[0], row[1], row[2]
        if sku and permalink:
            sku_to_permalink[str(sku).strip().upper()] = permalink
    return sku_to_permalink


def resolve(v_lista_precios_path, elice_urls_path):
    sku_to_permalink = load_elice_sku_index(elice_urls_path)

    wb = openpyxl.load_workbook(v_lista_precios_path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(min_row=1, max_row=1, values_only=True)
    header = list(next(rows_iter))

    resolved = []
    tier_counts = Counter()
    seen_stock = set()

    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(header, row))
        stock = (record.get("Stock") or "").strip()
        if not stock or stock in seen_stock:
            continue
        seen_stock.add(stock)

        cod_grupo_unidad = (record.get("CodGrupoUnidad") or "").strip().upper()

        tier1_permalink = sku_to_permalink.get(stock.upper())
        tier2_permalink = None
        tier2_group_code = None
        tier3_permalink = None
        tier3_group_code = None

        if not tier1_permalink and cod_grupo_unidad:
            tier2_permalink = sku_to_permalink.get(cod_grupo_unidad)
            tier2_group_code = cod_grupo_unidad

        if not tier1_permalink and not tier2_permalink and cod_grupo_unidad:
            normalized = normalize_group(cod_grupo_unidad)
            if normalized != cod_grupo_unidad:
                tier3_permalink = sku_to_permalink.get(normalized)
                tier3_group_code = normalized

        if tier1_permalink:
            match_by = "stock"
            permalink = tier1_permalink
            resolved_group_code = None
        elif tier2_permalink:
            match_by = "grupo"
            permalink = tier2_permalink
            resolved_group_code = tier2_group_code
        elif tier3_permalink:
            match_by = "grupo"
            permalink = tier3_permalink
            resolved_group_code = tier3_group_code
        else:
            match_by = "sin_match"
            permalink = ""
            resolved_group_code = None

        tier_counts[match_by] += 1

        resolved.append(
            {
                "stock": stock,
                "marca": (record.get("Marca") or "").strip(),
                "modelo": (record.get("Modelo") or "").strip(),
                "codGrupoUnidad": cod_grupo_unidad,
                "config": (record.get("Config.") or "").strip(),
                "susp": (record.get("Susp") or "").strip(),
                "tipoMotor": (record.get("Tipo Motor") or "").strip(),
                "permalink": permalink,
                "matchBy": match_by,
                "resolvedGroupCode": resolved_group_code,
            }
        )

    return resolved, tier_counts


def main():
    if len(sys.argv) != 4:
        print(
            "Uso: python resolve-permalinks.py <V_ListaPreciosTableau2.xlsx> "
            "<ELICE_URLS.xlsx> <salida.json>"
        )
        sys.exit(1)

    v_path, elice_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
    resolved, tier_counts = resolve(v_path, elice_path)

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(resolved, f, ensure_ascii=False, indent=2)

    print(f"Total filas: {len(resolved)}")
    print(f"Por Stock (foto propia): {tier_counts.get('stock', 0)}")
    print(f"Por CodGrupoUnidad (default): {tier_counts.get('grupo', 0)}")
    print(f"Sin match: {tier_counts.get('sin_match', 0)}")
    print(f"Escrito: {out_path}")


if __name__ == "__main__":
    main()
